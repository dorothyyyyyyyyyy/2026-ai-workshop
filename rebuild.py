"""
Rebuild slide to exactly match user's reference screenshot.
Creates a single-slide PPTX with:
- "Just Headline | White" layout
- Title: "Agenda - Wednesday 5/27/2026 day 1"
- 6-column, 22-row agenda table
- Number Tag Honey shapes for sub-items
"""
from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.oxml.ns import qn
from copy import deepcopy
import lxml.etree as ET
import uuid

SRC = "/Users/trip/Downloads/Test.pptx"
DST = "/Users/trip/Desktop/workshop/Test_formatted.pptx"

prs = Presentation(SRC)

# ── Find target layout ──
target_layout = None
for master in prs.slide_masters:
    for layout in master.slide_layouts:
        if layout.name == "Just Headline | White":
            target_layout = layout
            break
    if target_layout:
        break

# ── Get Number Tag template from slide 4 ──
slide4 = prs.slides[3]
tag_template_xml = None
for shape in slide4.shapes:
    if "Number Tag" in shape.name:
        tag_template_xml = deepcopy(shape._element)
        break

# ── Create a new slide using the target layout ──
new_slide = prs.slides.add_slide(target_layout)

# ── Set title ──
for shape in new_slide.shapes:
    if shape.has_text_frame:
        nvSpPr = shape._element.find(qn("p:nvSpPr"))
        if nvSpPr is not None:
            nvPr = nvSpPr.find(qn("p:nvPr"))
            if nvPr is not None:
                ph = nvPr.find(qn("p:ph"))
                if ph is not None and ph.get("type") == "title":
                    tf = shape.text_frame
                    tf.clear()
                    p = tf.paragraphs[0]
                    run = p.add_run()
                    run.text = "Agenda - Wednesday 5/27/2026 day 1"
                    rPr = run._r.find(qn("a:rPr"))
                    if rPr is None:
                        rPr = ET.SubElement(run._r, qn("a:rPr"))
                    rPr.set("lang", "en-US")
                    rPr.set("b", "1")
                    rPr.set("dirty", "0")
                    sf = ET.SubElement(rPr, qn("a:solidFill"))
                    ET.SubElement(sf, qn("a:schemeClr")).set("val", "dk1")
                    latin = ET.SubElement(rPr, qn("a:latin"))
                    latin.set("typeface", "Noto Sans")
                    latin.set("panose", "020B0502040504020204")
                    latin.set("pitchFamily", "34")
                    latin.set("charset", "0")

# ── Table data ──
# Row types: "header", "category", "sub"
# category rows: (section_num, item, time, holder, participants, description)
# sub rows:      (tag_num, item, time, holder, participants, description)

ROWS = [
    ("header", ["LEGEND PLACEHOLDER", "", "TIME", "HOLDER", "PARTICIPANTS", "DESCRIPTION"]),
    ("category", ["1", "Warm up", "00:00~09:35", "", "", "ROOM：TBD"]),
    ("sub", ["1", "Get together in the Meeting Room", "00:00~09:00", "", "", "Waking up"]),
    ("sub", ["2", "Welcome & Introduction", "09:00~09:30", "Simmon", "WS Members", "Greeting/netizen F2F meeting"]),
    ("sub", ["3", "Agenda introduction", "09:30~09:35", "Ziwen", "WS Members", "Telling you when you can take a break"]),
    ("category", ["2", "BU Strategy", "09:35~12:00", "", "", ""]),
    ("sub", ["1", "Strategy, MYP, 2026 focus", "09:35~10:35", "Lang", "Extension", ""]),
    ("sub", ["2", "Q&A", "10:35~10:45", "", "Extension", ""]),
    ("sub", ["3", "Break", "10:45~11:15", "", "", ""]),
    ("sub", ["4", "Market Overview", "11:15~11:45", "Ziwen", "Extension", ""]),
    ("sub", ["5", "Q&A and open discussion", "11:45~12:00", "", "Extension", ""]),
    ("category", ["3", "Lunch", "12:00~13:15", "", "", ""]),
    ("category", ["4", "R&D", "13:15~15:15", "", "", ""]),
    ("sub", ["1", "Dual-stage OD for Tsingshan EDS1-H", "13:15~13:40", "Lora", "Extension", "DFM Review based on Physical prototype + slides & design file"]),
    ("sub", ["2", "Standalone GaN DCDC HV/12/48", "13:40~14:00", "Lora", "Extension", "DFM Review based on Physical prototype + slides & design file"]),
    ("sub", ["3", "Single-stage GaN OBC+DCDC 7kw for China", "14:00~14:25", "Lora", "Extension", "DFM Review based on Physical prototype + slides & design file"]),
    ("sub", ["4", "Single-stage GaN OBC+DCDC 22kw/11kw for BMW", "14:25~14:50", "Lora", "Extension", "DFM Review based on Physical prototype + slides & design file"]),
    ("sub", ["5", "Q&A", "14:50~15:00", "", "Extension", ""]),
    ("sub", ["6", "Break", "15:00~15:15", "", "", ""]),
    ("category", ["5", "Line tour & Line concept", "15:15~16:30", "Eddie", "Stakeholders", ""]),
    ("category", ["6", "Open discussion/Q&A/Email checking", "16:30~18:00", "", "", ""]),
    ("category", ["7", "Team dinner", "18:00~23:59", "", "", ""]),
]

NUM_ROWS = len(ROWS)
NUM_COLS = 6

# Column widths
COL_W = [360000, 3852000, 1400000, 900000, 1300000, 3420000]
TABLE_W = sum(COL_W)
TABLE_X = 479375
TABLE_Y = 1338264
HEADER_H = 216000
ROW_H = 222000  # data rows — compact to fit within slide bounds

# ── Build table XML ──
nsmap = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}

def make_run_xml(text, sz=1200, bold=False, color_hex=None, cap=None, spc=None, font="Noto Sans"):
    r = ET.SubElement(ET.Element("dummy"), qn("a:r"))
    rPr = ET.SubElement(r, qn("a:rPr"))
    rPr.set("lang", "en-US")
    rPr.set("altLang", "zh-CN")
    rPr.set("sz", str(sz))
    rPr.set("dirty", "0")
    if bold:
        rPr.set("b", "1")
    if cap:
        rPr.set("cap", cap)
    if spc is not None:
        rPr.set("spc", str(spc))
    if color_hex:
        sf = ET.SubElement(rPr, qn("a:solidFill"))
        ET.SubElement(sf, qn("a:srgbClr")).set("val", color_hex)
    latin = ET.SubElement(rPr, qn("a:latin"))
    latin.set("typeface", font)
    latin.set("panose", "020B0502040504020204" if "Black" not in font else "020B0A02040504020204")
    latin.set("pitchFamily", "34")
    latin.set("charset", "0")
    ea = ET.SubElement(rPr, qn("a:ea"))
    ea.set("typeface", "+mn-ea")
    cs = ET.SubElement(rPr, qn("a:cs"))
    cs.set("typeface", font)
    cs.set("panose", "020B0502040504020204" if "Black" not in font else "020B0A02040504020204")
    cs.set("pitchFamily", "34")
    cs.set("charset", "0")
    t = ET.SubElement(r, qn("a:t"))
    t.text = text
    return r

def make_para(text, sz=1200, bold=False, color_hex=None, cap=None, spc=None, font="Noto Sans"):
    p = ET.Element(qn("a:p"))
    pPr = ET.SubElement(p, qn("a:pPr"))
    pPr.set("marL", "0")
    pPr.set("algn", "l")
    lnSpc = ET.SubElement(pPr, qn("a:lnSpc"))
    spcPct = ET.SubElement(lnSpc, qn("a:spcPct"))
    spcPct.set("val", "120000")
    spcBef = ET.SubElement(pPr, qn("a:spcBef"))
    ET.SubElement(spcBef, qn("a:spcPts")).set("val", "0")
    spcAft = ET.SubElement(pPr, qn("a:spcAft"))
    ET.SubElement(spcAft, qn("a:spcPts")).set("val", "0")
    ET.SubElement(pPr, qn("a:buNone"))
    if text:
        r = make_run_xml(text, sz=sz, bold=bold, color_hex=color_hex, cap=cap, spc=spc, font=font)
        p.append(r)
    endRPr = ET.SubElement(p, qn("a:endParaRPr"))
    endRPr.set("lang", "zh-CN")
    endRPr.set("altLang", "en-US")
    endRPr.set("sz", str(sz))
    endRPr.set("dirty", "0")
    if bold:
        endRPr.set("b", "1")
    if cap:
        endRPr.set("cap", cap)
    if spc is not None:
        endRPr.set("spc", str(spc))
    if color_hex:
        sf2 = ET.SubElement(endRPr, qn("a:solidFill"))
        ET.SubElement(sf2, qn("a:srgbClr")).set("val", color_hex)
    return p

def make_border(tag, fill_type="noFill", w="6350", color_hex=None, scheme_clr=None):
    el = ET.Element(qn(f"a:{tag}"))
    el.set("w", w)
    el.set("cap", "flat")
    el.set("cmpd", "sng")
    el.set("algn", "ctr")
    if fill_type == "noFill":
        ET.SubElement(el, qn("a:noFill"))
    elif fill_type == "solid":
        sf = ET.SubElement(el, qn("a:solidFill"))
        if scheme_clr:
            ET.SubElement(sf, qn("a:schemeClr")).set("val", scheme_clr)
        elif color_hex:
            ET.SubElement(sf, qn("a:srgbClr")).set("val", color_hex)
    ET.SubElement(el, qn("a:prstDash")).set("val", "solid")
    ET.SubElement(el, qn("a:round"))
    he = ET.SubElement(el, qn("a:headEnd"))
    he.set("type", "none"); he.set("w", "med"); he.set("len", "med")
    te = ET.SubElement(el, qn("a:tailEnd"))
    te.set("type", "none"); te.set("w", "med"); te.set("len", "med")
    return el

def make_cell(text, row_type, col_idx, is_category_num=False, merge_span=None):
    tc = ET.Element(qn("a:tc"))
    if merge_span:
        tc.set("gridSpan", str(merge_span))

    txBody = ET.SubElement(tc, qn("a:txBody"))
    ET.SubElement(txBody, qn("a:bodyPr"))
    ET.SubElement(txBody, qn("a:lstStyle"))

    if row_type == "header":
        p = make_para(text, sz=800, color_hex="878787", cap="all", spc=50)
    elif is_category_num:
        p = make_para(text, sz=2400, bold=True, font="Noto Sans Black", color_hex="237864")
    elif row_type == "category":
        is_time = (col_idx == 2)
        p = make_para(text, sz=1200, bold=True)
    else:
        p = make_para(text, sz=1200)
    txBody.append(p)

    # tcPr
    tcPr = ET.SubElement(tc, qn("a:tcPr"))
    if row_type == "header":
        tcPr.set("marL", "72000"); tcPr.set("marR", "72000")
        tcPr.set("marT", "0"); tcPr.set("marB", "72000")
        tcPr.set("anchor", "ctr")
    elif is_category_num:
        tcPr.set("marL", "72000"); tcPr.set("marR", "72000")
        tcPr.set("marT", "36000"); tcPr.set("marB", "36000")
        tcPr.set("anchor", "ctr")
    else:
        tcPr.set("marL", "72000"); tcPr.set("marR", "72000")
        tcPr.set("marT", "72000"); tcPr.set("marB", "72000")
        tcPr.set("anchor", "ctr")

    # Borders
    tcPr.append(make_border("lnL", "noFill"))
    tcPr.append(make_border("lnR", "noFill"))

    if row_type == "header":
        tcPr.append(make_border("lnT", "noFill"))
        tcPr.append(make_border("lnB", "solid", w="12700", scheme_clr="tx2"))
    elif row_type == "category":
        tcPr.append(make_border("lnT", "solid", w="12700", scheme_clr="tx2"))
        tcPr.append(make_border("lnB", "noFill"))
    else:
        tcPr.append(make_border("lnT", "solid", w="6350", color_hex="E3E3E3"))
        tcPr.append(make_border("lnB", "solid", w="6350", color_hex="E3E3E3"))

    ET.SubElement(tcPr, qn("a:noFill"))
    return tc

def make_merged_placeholder():
    tc = ET.Element(qn("a:tc"))
    tc.set("hMerge", "1")
    txBody = ET.SubElement(tc, qn("a:txBody"))
    ET.SubElement(txBody, qn("a:bodyPr"))
    ET.SubElement(txBody, qn("a:lstStyle"))
    p = ET.SubElement(txBody, qn("a:p"))
    ET.SubElement(p, qn("a:endParaRPr")).set("lang", "zh-CN")
    tcPr = ET.SubElement(tc, qn("a:tcPr"))
    return tc

# Build tbl element
tbl = ET.Element(qn("a:tbl"))
tblPr = ET.SubElement(tbl, qn("a:tblPr"))
tblPr.set("firstRow", "1")
tblPr.set("firstCol", "1")
tblPr.set("lastRow", "1")
tblPr.set("bandRow", "1")
styleId = ET.SubElement(tblPr, qn("a:tableStyleId"))
styleId.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"

tblGrid = ET.SubElement(tbl, qn("a:tblGrid"))
for w in COL_W:
    gc = ET.SubElement(tblGrid, qn("a:gridCol"))
    gc.set("w", str(w))

# Track row Y positions for tag placement
row_y_positions = []
current_y = TABLE_Y

for ri, (row_type, cells) in enumerate(ROWS):
    h = HEADER_H if row_type == "header" else ROW_H
    row_y_positions.append(current_y)
    current_y += h

    tr = ET.SubElement(tbl, qn("a:tr"))
    tr.set("h", str(h))

    if row_type == "header":
        # cols 0+1 merged as "LEGEND PLACEHOLDER"
        tc0 = make_cell(cells[0], "header", 0, merge_span=2)
        tr.append(tc0)
        tr.append(make_merged_placeholder())
        for ci in range(2, NUM_COLS):
            tc = make_cell(cells[ci], "header", ci)
            tr.append(tc)
    else:
        for ci in range(NUM_COLS):
            text = cells[ci]
            is_cat_num = (row_type == "category" and ci == 0)
            if row_type == "sub" and ci == 0:
                # Sub-item col 0: empty (tag shape goes on top)
                tc = make_cell("", row_type, ci)
            else:
                tc = make_cell(text, row_type, ci, is_category_num=is_cat_num)
            tr.append(tc)

TABLE_H = current_y - TABLE_Y

# ── Insert table as graphicFrame ──
spTree = new_slide._element.find(qn("p:cSld") + "/" + qn("p:spTree"))

gf = ET.SubElement(spTree, qn("p:graphicFrame"))
nvGfPr = ET.SubElement(gf, qn("p:nvGraphicFramePr"))
cNvPr = ET.SubElement(nvGfPr, qn("p:cNvPr"))
cNvPr.set("id", "100")
cNvPr.set("name", "Table plain")
cNvGfPr = ET.SubElement(nvGfPr, qn("p:cNvGraphicFramePr"))
grpLocks = ET.SubElement(cNvGfPr, qn("a:graphicFrameLocks"))
grpLocks.set("noGrp", "1")
ET.SubElement(nvGfPr, qn("p:nvPr"))

xfrm = ET.SubElement(gf, qn("p:xfrm"))
off = ET.SubElement(xfrm, qn("a:off"))
off.set("x", str(TABLE_X)); off.set("y", str(TABLE_Y))
ext = ET.SubElement(xfrm, qn("a:ext"))
ext.set("cx", str(TABLE_W)); ext.set("cy", str(TABLE_H))

graphic = ET.SubElement(gf, qn("a:graphic"))
graphicData = ET.SubElement(graphic, qn("a:graphicData"))
graphicData.set("uri", "http://schemas.openxmlformats.org/drawingml/2006/table")
graphicData.append(tbl)

# ── Add Number Tag shapes for sub-items ──
def make_number_tag(num_text, x, y):
    sp = deepcopy(tag_template_xml)
    # Update position
    xfrm = sp.find(qn("p:spPr") + "/" + qn("a:xfrm"))
    xfrm.find(qn("a:off")).set("x", str(x))
    xfrm.find(qn("a:off")).set("y", str(y))
    # Update text
    for t_el in sp.findall(".//" + qn("a:t")):
        t_el.text = num_text
    # Update id and name to be unique
    cNvPr = sp.find(qn("p:nvSpPr") + "/" + qn("p:cNvPr"))
    cNvPr.set("id", str(200 + int(num_text) + hash(str(y)) % 1000))
    # Remove relationship references that may not exist in new slide
    nvPr = sp.find(qn("p:nvSpPr") + "/" + qn("p:nvPr"))
    custData = nvPr.find(qn("p:custDataLst"))
    if custData is not None:
        nvPr.remove(custData)
    # Generate unique creationId
    extLst = cNvPr.find(qn("a:extLst"))
    if extLst is not None:
        for ext in extLst.findall(qn("a:ext")):
            for child in ext:
                if "creationId" in child.tag:
                    child.set("id", "{" + str(uuid.uuid4()).upper() + "}")
    return sp

# Tag positions: center of col 0, vertically centered in each sub-item row
tag_x = TABLE_X + (COL_W[0] - 216000) // 2  # center in col 0

sub_item_rows = []
for ri, (row_type, cells) in enumerate(ROWS):
    if row_type == "sub":
        sub_item_rows.append((ri, cells[0]))

shape_id_counter = 200
for ri, tag_num in sub_item_rows:
    row_y = row_y_positions[ri]
    tag_y = row_y + (ROW_H - 216000) // 2
    tag_sp = make_number_tag(tag_num, tag_x, tag_y)
    tag_sp.find(qn("p:nvSpPr") + "/" + qn("p:cNvPr")).set("id", str(shape_id_counter))
    shape_id_counter += 1
    spTree.append(tag_sp)

# ── Now delete the original slides 1-3 and move the new slide to position ──
# The new slide was added at the end (index 4). We need to:
# 1. Remove slides 0, 1, 2 (original slides 1-3)
# 2. Keep slide 3 (original slide 4 as reference)
# 3. Keep the new slide (index 4)
# Actually: user wants just this one slide reproduced. Let's keep only the new slide.

# Remove slides 0-3 (original 4 slides), keeping only the new one
slide_id_list = prs.slides._sldIdLst
slide_ids = list(slide_id_list)
for sldId in slide_ids[:4]:  # remove original 4 slides
    rId = sldId.get(qn("r:id"))
    prs.part.drop_rel(rId)
    slide_id_list.remove(sldId)

prs.save(DST)
print(f"Saved to {DST}")
print(f"Slide count: {len(prs.slides)}")
print(f"Table rows: {NUM_ROWS}, cols: {NUM_COLS}")
print(f"Number tags added: {len(sub_item_rows)}")
