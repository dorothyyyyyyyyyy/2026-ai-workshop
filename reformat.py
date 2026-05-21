from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.oxml.ns import qn
from copy import deepcopy
import lxml.etree as ET

SRC = "/Users/trip/Downloads/Test.pptx"
DST = "/Users/trip/Desktop/workshop/Test_formatted.pptx"

prs = Presentation(SRC)

# --- Find "Just Headline | White" layout from Master 1 ---
target_layout = None
for master in prs.slide_masters:
    for layout in master.slide_layouts:
        if layout.name == "Just Headline | White":
            target_layout = layout
            break
    if target_layout:
        break
assert target_layout, "Could not find 'Just Headline | White' layout"

# --- Helper: build border XML ---
def make_border_noFill(tag, w="12700"):
    el = ET.SubElement(ET.Element("dummy"), qn(f"a:{tag}"))
    el.set("w", w)
    el.set("cap", "flat")
    el.set("cmpd", "sng")
    el.set("algn", "ctr")
    nf = ET.SubElement(el, qn("a:noFill"))
    ET.SubElement(el, qn("a:prstDash")).set("val", "solid")
    ET.SubElement(el, qn("a:round"))
    he = ET.SubElement(el, qn("a:headEnd"))
    he.set("type", "none"); he.set("w", "med"); he.set("len", "med")
    te = ET.SubElement(el, qn("a:tailEnd"))
    te.set("type", "none"); te.set("w", "med"); te.set("len", "med")
    return el

def make_border_solid(tag, w="6350", color_hex=None, scheme_clr=None):
    el = ET.SubElement(ET.Element("dummy"), qn(f"a:{tag}"))
    el.set("w", w)
    el.set("cap", "flat")
    el.set("cmpd", "sng")
    el.set("algn", "ctr")
    sf = ET.SubElement(el, qn("a:solidFill"))
    if scheme_clr:
        ET.SubElement(sf, qn("a:schemeClr")).set("val", scheme_clr)
    elif color_hex:
        ET.SubElement(sf, qn("a:srgbClr")).set("val", color_hex)
    ET.SubElement(el, qn("a:prstDash")).set("val", "solid")
    ET.SubElement(el, qn("a:round"))
    he = ET.SubElement(el, qn("a:headEnd"))
    he.set("type", "none"); he.set("w", "med"); he.set("len", "med")
    te = ET.SubElement(el, qn("a:tailEnd"))
    te.set("type", "none"); te.set("w", "med"); te.set("len", "med")
    return el

def set_cell_borders(tcPr, lnL=None, lnR=None, lnT=None, lnB=None):
    for tag in ["lnL", "lnR", "lnT", "lnB", "lnTlToBr", "lnBlToTr"]:
        old = tcPr.find(qn(f"a:{tag}"))
        if old is not None:
            tcPr.remove(old)
    if lnL is not None: tcPr.insert(0, lnL)
    if lnR is not None: tcPr.insert(1, lnR)
    if lnT is not None: tcPr.insert(2, lnT)
    if lnB is not None: tcPr.insert(3, lnB)

def set_run_font(rPr, font="Noto Sans", sz=None, color_hex=None, cap=None, spc=None, bold=None):
    if sz is not None:
        rPr.set("sz", str(sz))
    if cap is not None:
        rPr.set("cap", cap)
    if spc is not None:
        rPr.set("spc", str(spc))
    if bold is not None:
        rPr.set("b", "1" if bold else "0")
    # Set font
    for tag in ["a:latin", "a:cs"]:
        old = rPr.find(qn(tag))
        if old is not None:
            rPr.remove(old)
    latin = ET.SubElement(rPr, qn("a:latin"))
    latin.set("typeface", font)
    latin.set("panose", "020B0502040504020204")
    latin.set("pitchFamily", "34")
    latin.set("charset", "0")
    cs = ET.SubElement(rPr, qn("a:cs"))
    cs.set("typeface", font)
    cs.set("panose", "020B0502040504020204")
    cs.set("pitchFamily", "34")
    cs.set("charset", "0")
    # Set color
    if color_hex:
        old_fill = rPr.find(qn("a:solidFill"))
        if old_fill is not None:
            rPr.remove(old_fill)
        sf = ET.SubElement(rPr, qn("a:solidFill"))
        ET.SubElement(sf, qn("a:srgbClr")).set("val", color_hex)

def style_cell_text(cell_tc, font="Noto Sans", sz=1200, color_hex=None, cap=None, spc=None, bold=None, is_header=False):
    for p in cell_tc.findall(qn("a:txBody") + "/" + qn("a:p")):
        # Style paragraph properties
        pPr = p.find(qn("a:pPr"))
        if pPr is None:
            pPr = ET.SubElement(p, qn("a:pPr"))
        # Set line spacing 120%
        lnSpc = pPr.find(qn("a:lnSpc"))
        if lnSpc is None:
            lnSpc = ET.SubElement(pPr, qn("a:lnSpc"))
        else:
            for ch in list(lnSpc):
                lnSpc.remove(ch)
        spcPct = ET.SubElement(lnSpc, qn("a:spcPct"))
        spcPct.set("val", "120000")

        for r in p.findall(qn("a:r")):
            rPr = r.find(qn("a:rPr"))
            if rPr is None:
                rPr = ET.SubElement(r, qn("a:rPr"))
            set_run_font(rPr, font=font, sz=sz, color_hex=color_hex, cap=cap, spc=spc, bold=bold)

        endParaRPr = p.find(qn("a:endParaRPr"))
        if endParaRPr is not None:
            set_run_font(endParaRPr, font=font, sz=sz, color_hex=color_hex, cap=cap, spc=spc, bold=bold)

def update_tcPr(cell_tc, marL=72000, marR=72000, marT=72000, marB=72000, anchor="ctr"):
    tcPr = cell_tc.find(qn("a:tcPr"))
    if tcPr is None:
        tcPr = ET.SubElement(cell_tc, qn("a:tcPr"))
    tcPr.set("marL", str(marL))
    tcPr.set("marR", str(marR))
    tcPr.set("marT", str(marT))
    tcPr.set("marB", str(marB))
    tcPr.set("anchor", anchor)
    # Remove old fill, set noFill
    for tag in ["a:solidFill", "a:noFill", "a:gradFill", "a:pattFill"]:
        old = tcPr.find(qn(tag))
        if old is not None:
            tcPr.remove(old)
    ET.SubElement(tcPr, qn("a:noFill"))
    return tcPr


# --- Process slides 1-3 ---
for slide_idx in range(3):
    slide = prs.slides[slide_idx]

    # 1. Change layout — remove old relationship first to avoid duplicates
    slide._element.attrib.clear()
    rels = slide.part._rels
    old_rIds = [rId for rId, rel in rels.items()
                if rel.reltype == "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideLayout"]
    for old_rId in old_rIds:
        rels.pop(old_rId)
    rId = slide.part.relate_to(target_layout.part, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideLayout")

    # 2. Remove unwanted shapes (think-cell OLE, floating textbox on slide 2)
    spTree = slide._element.find(qn("p:cSld") + "/" + qn("p:spTree"))
    shapes_to_remove = []
    for child in list(spTree):
        # Remove OLE objects (think-cell)
        if child.tag == qn("p:graphicFrame"):
            nvPr = child.find(qn("p:nvGraphicFramePr"))
            if nvPr is not None:
                cNvPr = nvPr.find(qn("p:cNvPr"))
                if cNvPr is not None and "think-cell" in (cNvPr.get("name") or ""):
                    shapes_to_remove.append(child)
        # Remove floating textbox on slide 2
        if slide_idx == 1 and child.tag == qn("p:sp"):
            nvSpPr = child.find(qn("p:nvSpPr"))
            if nvSpPr is not None:
                cNvPr = nvSpPr.find(qn("p:cNvPr"))
                if cNvPr is not None and cNvPr.get("name") == "TextBox 6":
                    shapes_to_remove.append(child)

    for shape_el in shapes_to_remove:
        spTree.remove(shape_el)

    # 3. Reposition title and subtitle placeholders
    for sp in spTree.findall(qn("p:sp")):
        nvSpPr = sp.find(qn("p:nvSpPr"))
        if nvSpPr is None:
            continue
        nvPr = nvSpPr.find(qn("p:nvPr"))
        if nvPr is None:
            continue
        ph = nvPr.find(qn("p:ph"))
        if ph is None:
            continue

        ph_type = ph.get("type")
        ph_idx = ph.get("idx")
        cNvPr = nvSpPr.find(qn("p:cNvPr"))
        name = cNvPr.get("name") if cNvPr is not None else ""

        spPr = sp.find(qn("p:spPr"))

        if name == "Titel 8" or ph_type == "title":
            # Title placeholder -> match slide 4 position
            xfrm = spPr.find(qn("a:xfrm"))
            if xfrm is None:
                xfrm = ET.SubElement(spPr, qn("a:xfrm"))
            off = xfrm.find(qn("a:off"))
            if off is None:
                off = ET.SubElement(xfrm, qn("a:off"))
            off.set("x", "479375")
            off.set("y", "720000")
            ext = xfrm.find(qn("a:ext"))
            if ext is None:
                ext = ET.SubElement(xfrm, qn("a:ext"))
            ext.set("cx", "8928625")
            ext.set("cy", "246221")

            # Apply Noto Sans to title text
            txBody = sp.find(qn("p:txBody"))
            if txBody is not None:
                for p_el in txBody.findall(qn("a:p")):
                    for r in p_el.findall(qn("a:r")):
                        rPr = r.find(qn("a:rPr"))
                        if rPr is None:
                            rPr = ET.SubElement(r, qn("a:rPr"))
                        # Set Noto Sans font
                        for tag in ["a:latin", "a:cs"]:
                            old = rPr.find(qn(tag))
                            if old is not None:
                                rPr.remove(old)
                        latin = ET.SubElement(rPr, qn("a:latin"))
                        latin.set("typeface", "Noto Sans")
                        latin.set("panose", "020B0502040504020204")
                        latin.set("pitchFamily", "34")
                        latin.set("charset", "0")
                        # Set color to dk1
                        old_fill = rPr.find(qn("a:solidFill"))
                        if old_fill is not None:
                            rPr.remove(old_fill)
                        sf = ET.SubElement(rPr, qn("a:solidFill"))
                        ET.SubElement(sf, qn("a:schemeClr")).set("val", "dk1")

        elif name == "Textplatzhalter 10" or (ph_type == "body" and ph_idx == "13"):
            # Subtitle placeholder
            xfrm = spPr.find(qn("a:xfrm"))
            if xfrm is None:
                xfrm = ET.SubElement(spPr, qn("a:xfrm"))
            off = xfrm.find(qn("a:off"))
            if off is None:
                off = ET.SubElement(xfrm, qn("a:off"))
            off.set("x", "479375")
            off.set("y", "476024")
            ext = xfrm.find(qn("a:ext"))
            if ext is None:
                ext = ET.SubElement(xfrm, qn("a:ext"))
            ext.set("cx", "8928625")
            ext.set("cy", "115416")

    # 4. Style the table
    for gf in spTree.findall(qn("p:graphicFrame")):
        nvPr = gf.find(qn("p:nvGraphicFramePr"))
        if nvPr is None:
            continue
        cNvPr = nvPr.find(qn("p:cNvPr"))
        if cNvPr is not None and "think-cell" in (cNvPr.get("name") or ""):
            continue

        graphic = gf.find(qn("a:graphic"))
        if graphic is None:
            continue
        graphicData = graphic.find(qn("a:graphicData"))
        if graphicData is None or graphicData.get("uri") != "http://schemas.openxmlformats.org/drawingml/2006/table":
            continue

        tbl = graphicData.find(qn("a:tbl"))
        if tbl is None:
            continue

        # Reposition table
        xfrm = gf.find(qn("p:xfrm"))
        if xfrm is None:
            xfrm = ET.SubElement(gf, qn("p:xfrm"))
        off = xfrm.find(qn("a:off"))
        if off is None:
            off = ET.SubElement(xfrm, qn("a:off"))
        ext = xfrm.find(qn("a:ext"))
        if ext is None:
            ext = ET.SubElement(xfrm, qn("a:ext"))

        old_width = int(ext.get("cx", "11667014"))
        off.set("x", "479375")
        off.set("y", "1338264")
        new_width = 11232000
        ext.set("cx", str(new_width))
        # Keep height as-is (content-dependent)

        # Update table style ID
        tblPr = tbl.find(qn("a:tblPr"))
        if tblPr is not None:
            tblPr.set("firstRow", "1")
            tblPr.set("firstCol", "1")
            tblPr.set("lastRow", "1")
            tblPr.set("bandRow", "1")
            styleId = tblPr.find(qn("a:tableStyleId"))
            if styleId is not None:
                styleId.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"

        # Redistribute column widths proportionally
        tblGrid = tbl.find(qn("a:tblGrid"))
        if tblGrid is not None:
            cols = tblGrid.findall(qn("a:gridCol"))
            total_old = sum(int(c.get("w", "0")) for c in cols)
            if total_old > 0:
                for c in cols:
                    old_w = int(c.get("w", "0"))
                    new_w = int(old_w * new_width / total_old)
                    c.set("w", str(new_w))

        # Style each row/cell
        rows = tbl.findall(qn("a:tr"))
        for ri, tr in enumerate(rows):
            is_header = (ri == 0)
            cells = tr.findall(qn("a:tc"))

            for ci, tc in enumerate(cells):
                # Update cell margins
                tcPr = update_tcPr(tc, marL=72000, marR=72000,
                                   marT=0 if is_header else 72000,
                                   marB=72000)

                # Set borders
                lnL = make_border_noFill("lnL")
                lnR = make_border_noFill("lnR")

                if is_header:
                    lnT = make_border_noFill("lnT")
                    lnB = make_border_solid("lnB", w="12700", scheme_clr="tx2")
                else:
                    lnT = make_border_solid("lnT", w="6350", color_hex="E3E3E3")
                    lnB = make_border_solid("lnB", w="6350", color_hex="E3E3E3")

                set_cell_borders(tcPr, lnL, lnR, lnT, lnB)

                # Style text
                if is_header:
                    style_cell_text(tc, font="Noto Sans", sz=800, color_hex="878787",
                                    cap="all", spc=50, is_header=True)
                else:
                    style_cell_text(tc, font="Noto Sans", sz=1200)

    print(f"Slide {slide_idx+1} reformatted.")

prs.save(DST)
print(f"\nSaved to {DST}")
